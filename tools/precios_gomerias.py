"""Planilla de precios por cantidad para gomerías (exportes/PRECIOS GOMERIAS.xlsx).
Sin precios: el usuario los carga a mano (celdas amarillas) con lo que confirme el dueño."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

LLANTAS = ["Llanta aleación R14", "Llanta aleación R15", "Llanta aleación R16", "Llanta aleación R17",
           "Llanta de chapa (auto)", "Llanta de chapa (camioneta)", "Auxilio sin rodar"]
CUBIERTAS = ["Cubierta usada (todas las medidas)"]
ESC_LL = ["1 a 5", "6 a 10", "11 a 20", "21 a 50", "51 a 100"]
ESC_CU = ["1 a 5", "6 a 10", "11 a 20", "21 a 40", "41 a 80"]

F = "Arial"
NEG = PatternFill("solid", fgColor="16181B"); AMA = PatternFill("solid", fgColor="FFF2B3")
borde = Border(bottom=Side(style="thin", color="D0D0D0"), right=Side(style="thin", color="FFFFFF"))

def hoja(ws, titulo, filas, escalas):
    ws["A1"] = titulo; ws["A1"].font = Font(name=F, size=16, bold=True)
    ws["A2"] = "Precio por unidad según la cantidad que lleva la gomería. Completar las celdas amarillas."
    ws["A2"].font = Font(name=F, size=10, italic=True, color="666666")
    ws.cell(4, 1, "Producto")
    for j, e in enumerate(escalas):
        ws.cell(4, 2 + j, f"{e} unidades")
    for c in ws[4][:1 + len(escalas)]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = NEG
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 30
    for i, nombre in enumerate(filas):
        r = 5 + i
        a = ws.cell(r, 1, nombre); a.font = Font(name=F, bold=True); a.border = borde
        ws.row_dimensions[r].height = 22
        for j in range(len(escalas)):
            c = ws.cell(r, 2 + j); c.fill = AMA; c.border = borde
            c.number_format = '"$"#,##0'; c.font = Font(name=F, color="0000FF", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF": ws.column_dimensions[col].width = 16
    ws.freeze_panes = "B5"

wb = Workbook()
wb.active.title = "Llantas"
hoja(wb.active, "Llantas — precios por cantidad", LLANTAS, ESC_LL)
hoja(wb.create_sheet("Cubiertas"), "Cubiertas — precios por cantidad", CUBIERTAS, ESC_CU)
wb.save(sys.argv[1])
print(sys.argv[1])
